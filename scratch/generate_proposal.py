import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_premium_business_proposal():
    doc = Document()

    # Configure Margins (0.8 inch all around for expansive modern layout)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Tokens
    C_NAVY = RGBColor(11, 14, 20)          # #0B0E14 Dark Slate Navy
    C_GOLD = RGBColor(212, 160, 23)        # #D4A017 Tactical Gold
    C_CYAN = RGBColor(0, 180, 216)         # #00B4D8 Cyber Cyan
    C_CHARCOAL = RGBColor(33, 37, 45)      # #21252D Charcoal
    C_MUTED = RGBColor(108, 117, 125)      # #6C757D Muted Slate
    C_GREEN = RGBColor(46, 213, 115)       # #2ED573 Emerald Green

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_table_borders(table, border_hex="E2E8F0"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_hex}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_callout_box(title, text, bg_hex="F8FAFC", border_hex="D4A017"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.9)
        set_cell_background(cell, bg_hex)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.top_spacing = Pt(8)
        p.paragraph_format.bottom_spacing = Pt(8)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)

        r_title = p.add_run(f"{title}\n")
        r_title.font.name = "Arial"
        r_title.font.bold = True
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = C_GOLD

        r_body = p.add_run(text)
        r_body.font.name = "Calibri"
        r_body.font.size = Pt(10)
        r_body.font.color.rgb = C_CHARCOAL

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_screenshot_figure(img_path, caption):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_path, width=Inches(6.6))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(14)
            r_cap = p_cap.add_run(f"📸 {caption}")
            r_cap.font.name = "Calibri"
            r_cap.font.size = Pt(9)
            r_cap.font.italic = True
            r_cap.font.color.rgb = C_MUTED

    # ─────────────────────────────────────────────────────────────────────────────
    # COVER PAGE
    # ─────────────────────────────────────────────────────────────────────────────
    p_cov_top = doc.add_paragraph()
    p_cov_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cov_top.paragraph_format.space_before = Pt(30)
    p_cov_top.paragraph_format.space_after = Pt(10)

    emblem_path = "/Users/madushanminuwantha/Desktop/project/UWE/src/assets/images/uwe_hero_emblem.png"
    if os.path.exists(emblem_path):
        p_cov_top.add_run().add_picture(emblem_path, width=Inches(2.0))

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("UNITY WARRIORS EMPIRE\n")
    r_t1.font.name = "Arial"
    r_t1.font.size = Pt(26)
    r_t1.font.bold = True
    r_t1.font.color.rgb = C_GOLD

    r_t2 = p_title.add_run("COMMERCIAL BUSINESS PROPOSAL & TECHNICAL SYSTEM WHITE-PAPER")
    r_t2.font.name = "Arial"
    r_t2.font.size = Pt(14)
    r_t2.font.bold = True
    r_t2.font.color.rgb = C_NAVY
    p_title.paragraph_format.space_after = Pt(25)

    add_callout_box(
        "CONFIDENTIAL & PROPRIETARY COMMERCIAL PROPOSAL",
        "This document details the complete operational capabilities, dual-tier software architecture, "
        "automated student enrollment pipeline, and defense-grade security matrix of the Unity Warriors Empire "
        "training & academy platform. Ready for immediate enterprise operation, commercial licensing, or white-label acquisition."
    )

    p_cov_meta = doc.add_paragraph()
    p_cov_meta.paragraph_format.space_before = Pt(60)
    p_cov_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_cov_meta.add_run(
        "Author: UWE Systems Architecture & Engineering Division\n"
        "Document Release: Version 2.4 (Production Validated)\n"
        "Deployment Target: Vercel Cloud Serverless + PostgreSQL Supabase\n"
        "Quality Assurance: 32 / 32 Automated Test Suites Passing (100%)"
    )
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = C_MUTED

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────────────────────
    # SECTION 1: EXECUTIVE OVERVIEW & VALUE PROPOSITION
    # ─────────────────────────────────────────────────────────────────────────────
    h1 = doc.add_heading("1. Executive Overview & Commercial Value Proposition", level=1)
    h1.runs[0].font.color.rgb = C_NAVY

    doc.add_paragraph(
        "Modern high-ticket coaching, subconscious mind optimization, and executive leadership academies require "
        "a digital infrastructure that conveys elite operational excellence. Generic educational templates and bloated "
        "third-party platforms (such as Teachable or Kajabi) introduce heavy monthly recurring SaaS costs, high latency, "
        "and fragmented workflows."
    )

    doc.add_paragraph(
        "Unity Warriors Empire (UWE) was engineered from the ground up to solve these challenges by unifying a "
        "cinematic student-facing academy with a centralized 12-in-1 Command HQ."
    )

    add_callout_box(
        "KEY COMMERCIAL ADVANTAGES",
        "• 100% Owned IP & Zero Recurring Platform Fees: Self-hosted serverless architecture eliminates monthly SaaS licensing overhead.\n"
        "• Frictionless Bank Transfer Monetization: Custom canvas-compressed slip submission and 1-click admin auto-enrolment.\n"
        "• Integrated Student Command Center: Personalized progress tracking, live Zoom links, and official credential vaults.\n"
        "• Granular Access Protection: Public free previews paired with strict token-gated masterclass content."
    )

    # ─────────────────────────────────────────────────────────────────────────────
    # SECTION 2: VISUAL PROOFS & LIVE OPERATIONAL SHOWCASE
    # ─────────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    h2 = doc.add_heading("2. Visual Platform Showcase & Live Production Screenshots", level=1)
    h2.runs[0].font.color.rgb = C_NAVY

    doc.add_paragraph(
        "The screenshots below demonstrate the live, fully responsive user experience across both the "
        "student learning environment and the back-office administrative command suite:"
    )

    # Screenshot 1: Student Dashboard
    img_dashboard = "/Users/madushanminuwantha/.gemini/antigravity-ide/brain/24430a5e-8bce-40c3-8180-a9f99641d361/.user_uploaded/media_1787491341436.png"
    add_screenshot_figure(
        img_dashboard,
        "Figure 1: Operative Student Command Dashboard — Live Program Directives, KPI Progress Counters, and Direct Zoom Mastermind Launchpad"
    )

    # Screenshot 2: Tactical Video Showcase
    img_demos = "/Users/madushanminuwantha/.gemini/antigravity-ide/brain/24430a5e-8bce-40c3-8180-a9f99641d361/.user_uploaded/media_1787491315601.png"
    add_screenshot_figure(
        img_demos,
        "Figure 2: Tactical Video Masterclass Vault — High-Fidelity Custom Thumbnails, Video Duration Timers, and Interactive Modal Player"
    )

    # Screenshot 3: LMS Series CMS
    img_series = "/Users/madushanminuwantha/.gemini/antigravity-ide/brain/24430a5e-8bce-40c3-8180-a9f99641d361/.user_uploaded/media_1787491371847.png"
    add_screenshot_figure(
        img_series,
        "Figure 3: Admin Command HQ: Video Series & Gated LMS Vault — Drag-and-Drop Thumbnail Cover Uploader with Direct PostgreSQL Persistence"
    )

    # Screenshot 4: WhatsApp Leads CRM
    img_leads = "/Users/madushanminuwantha/.gemini/antigravity-ide/brain/24430a5e-8bce-40c3-8180-a9f99641d361/.user_uploaded/media_1787491396614.png"
    add_screenshot_figure(
        img_leads,
        "Figure 4: Admin Inbound Leads CRM — WhatsApp Ingestion Pipeline, Status Lifecycle Management (NEW, CONTACTED, ENROLLED), and 1-Click Operative Conversion"
    )

    # Screenshot 5: Recruitment & Capacity Sync
    img_jobs = "/Users/madushanminuwantha/.gemini/antigravity-ide/brain/24430a5e-8bce-40c3-8180-a9f99641d361/.user_uploaded/media_1787491414209.png"
    add_screenshot_figure(
        img_jobs,
        "Figure 5: Talent Acquisition & Live Capacity Sync — Visual Open Position Capacity Bars, Hiring Statuses, and Applicant Tracking"
    )

    # ─────────────────────────────────────────────────────────────────────────────
    # SECTION 3: COMPLETE FRONTEND FEATURE SPECIFICATION
    # ─────────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    h3 = doc.add_heading("3. Comprehensive Frontend Capabilities & User Experience", level=1)
    h3.runs[0].font.color.rgb = C_NAVY

    doc.add_paragraph(
        "Built using React 19 and Framer Motion with tailored dark-mode ergonomics, the frontend delivers "
        "sub-second responsiveness and high-ticket perceived value."
    )

    fe_table = doc.add_table(rows=1, cols=3)
    fe_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(fe_table)

    h_cells = fe_table.rows[0].cells
    h_cells[0].text = "Module / Area"
    h_cells[1].text = "Key Functional Features"
    h_cells[2].text = "Business / UX Impact"
    
    for c in h_cells:
        set_cell_background(c, "0B0E14")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = C_GOLD
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

    h_cells[0].width = Inches(1.5)
    h_cells[1].width = Inches(3.2)
    h_cells[2].width = Inches(2.2)

    fe_data = [
        ("Home Portal", "Spring-physics cursor lighting, tactical hero cards, dynamic stats ticker, and animated value propositions.", "Commands immediate attention and authority upon first page load."),
        ("Program Catalog", "Real-time pricing, seat scarcity countdowns, and upcoming batch dates fetched dynamically from Supabase.", "Drives conversion urgency with real-time seat availability."),
        ("Course Detail View", "Modular day-by-day syllabuses, faculty credentials, interactive FAQs, and verified student feedback.", "Addresses prospective student objections with transparent curriculum data."),
        ("Student Dashboard", "Personalized enrolled directives, module watch meters, schedule calendars, and direct Zoom launch buttons.", "Provides an organized central command center for paying students."),
        ("Gated LMS Video Vault", "Category-wise series playlists (BMB, Leadership, IGNIT) with public Episode 1 Free Previews and token gating.", "Nurtures free leads while strictly protecting premium intellectual property."),
        ("Bank Slip Uploader", "Client-side HTML5 Canvas image optimization (PNG, JPG, WebP) with live status tracking (PENDING, VERIFIED).", "Enables frictionless local payments with zero payment gateway fees."),
        ("Careers Portal", "Live open position listings with department filters, income ranges, and instant application forms.", "Automates recruitment and talent acquisition for internal growth."),
        ("Global SEO Engine", "Pre-configured Open Graph, canonical links, and dynamic meta descriptions via React Helmet Async.", "Maximizes organic search visibility across search engines and social sharing.")
    ]

    for idx, (m, f, b) in enumerate(fe_data):
        row = fe_table.add_row().cells
        row[0].text = m
        row[1].text = f
        row[2].text = b
        if idx % 2 == 1:
            set_cell_background(row[0], "F8FAFC")
            set_cell_background(row[1], "F8FAFC")
            set_cell_background(row[2], "F8FAFC")
        row[0].paragraphs[0].runs[0].font.bold = True
        row[0].paragraphs[0].runs[0].font.size = Pt(9)
        row[1].paragraphs[0].runs[0].font.size = Pt(9)
        row[2].paragraphs[0].runs[0].font.size = Pt(9)

    # ─────────────────────────────────────────────────────────────────────────────
    # SECTION 4: BACKEND & ADMIN COMMAND HQ SPECIFICATION
    # ─────────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    h4 = doc.add_heading("4. Centralized Admin Command HQ & Operations Matrix", level=1)
    h4.runs[0].font.color.rgb = C_NAVY

    doc.add_paragraph(
        "Command HQ serves as the complete operational cockpit for business owners, training coordinators, "
        "and recruitment staff. All changes synchronize with the PostgreSQL database in real time."
    )

    be_table = doc.add_table(rows=1, cols=3)
    be_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(be_table)

    bh_cells = be_table.rows[0].cells
    bh_cells[0].text = "HQ Control Tab"
    bh_cells[1].text = "Administrative Capability"
    bh_cells[2].text = "Operational Benefit"
    
    for c in bh_cells:
        set_cell_background(c, "0B0E14")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = C_GOLD
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

    bh_cells[0].width = Inches(1.6)
    bh_cells[1].width = Inches(3.2)
    bh_cells[2].width = Inches(2.1)

    be_data = [
        ("Batches, Fees & Zoom Links", "Edit tuition fees, available seats, commencement dates, and live Mastermind Zoom meeting URLs per batch.", "Instant scheduling updates reflected immediately on student portals."),
        ("Bank Slips Hub", "Visual inspection of uploaded payment receipts with 1-click verification that automatically enrolls students.", "Eliminates manual bookkeeping and accelerates student onboarding."),
        ("WhatsApp CRM Leads", "Live lead ingestion pipeline with statuses (NEW, CONTACTED, ENROLLED, REJECTED) and 1-click conversion to user accounts.", "Maximizes sales follow-up velocity and tracks marketing ROI."),
        ("Coupons & Promos", "Create discount vouchers with percentage or fixed deductions, max uses, expiry dates, and program restrictions.", "Empowers flexible promotional campaigns and affiliate marketing."),
        ("Video & Media CMS", "Publish and edit video reels and masterclasses with drag-and-drop thumbnail uploads and canvas compression.", "Zero developer dependencies for curriculum updates."),
        ("User Accounts Registry", "Manage student profiles, grant special course clearances, and perform instant security password resets.", "Complete governance over enrolled user records."),
        ("Ticker Announcement Bar", "Global emergency or promotional alert banner editor with dirty-state change detection.", "Instantly broadcast critical deadlines or final seat callouts."),
        ("Job Vacancies Sync", "Create and adjust open career slots with visual progress bars and applicant review tables.", "Streamlines hiring of sales specialists and coaches.")
    ]

    for idx, (m, f, b) in enumerate(be_data):
        row = be_table.add_row().cells
        row[0].text = m
        row[1].text = f
        row[2].text = b
        if idx % 2 == 1:
            set_cell_background(row[0], "F8FAFC")
            set_cell_background(row[1], "F8FAFC")
            set_cell_background(row[2], "F8FAFC")
        row[0].paragraphs[0].runs[0].font.bold = True
        row[0].paragraphs[0].runs[0].font.size = Pt(9)
        row[1].paragraphs[0].runs[0].font.size = Pt(9)
        row[2].paragraphs[0].runs[0].font.size = Pt(9)

    # ─────────────────────────────────────────────────────────────────────────────
    # SECTION 5: ENTERPRISE SECURITY & DEFENSE ARCHITECTURE
    # ─────────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    h5 = doc.add_heading("5. Enterprise Security & Defense Architecture", level=1)
    h5.runs[0].font.color.rgb = C_NAVY

    doc.add_paragraph(
        "Security has been built with defense-in-depth principles across every network boundary, "
        "authentication handshake, and database interaction."
    )

    sec_cards = [
        ("🔐 Dual-Token Authentication Protocol", "Short-lived 15-minute Access Tokens (JWT) paired with cryptographic 7-day Refresh Tokens. If an access token is intercepted, it expires rapidly, while refresh tokens are rotated safely."),
        ("🛡️ Granular Role-Based Access Control (RBAC)", "Hierarchical authorization enforcing four discrete clearance roles: SUPER_ADMIN, COMMANDER, COACH, and RECRUITER. Financial and credentialing endpoints strictly reject unauthorized staff."),
        ("🚦 3-Tier Layered Rate Limiting", "Dedicated express-rate-limiters: Auth (12 attempts/15min for brute-force defense), Submissions (20/hr for lead spam defense), and Global API (300/15min), with Vercel 'trust proxy' client IP identification."),
        ("🔍 Schema Enforcement & Payload Sanitization (Zod)", "Every API mutation passes through strict Zod validators, automatically stripping unexpected keys and rejecting malformed payloads before database execution."),
        ("🔑 Bcrypt Salted Password Hashing", "All administrative and student passwords are encrypted with Bcrypt using 12 salt rounds, ensuring protection against rainbow table attacks."),
        ("🗄️ PgBouncer Database Connection Pooling", "Supabase PostgreSQL transactions communicate via PgBouncer on port 6543, providing elasticity against serverless concurrency surges.")
    ]

    for title, desc in sec_cards:
        add_callout_box(title, desc, bg_hex="F8FAFC", border_hex="0B0E14")

    # ─────────────────────────────────────────────────────────────────────────────
    # SECTION 6: AUTOMATED MONETIZATION WORKFLOW & QA METRICS
    # ─────────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    h6 = doc.add_heading("6. Automated Student Lifecycle & Technical QA Metrics", level=1)
    h6.runs[0].font.color.rgb = C_NAVY

    doc.add_paragraph("The platform powers a continuous, automated 6-stage monetization pipeline:")

    flow_steps = [
        ("1. Inbound Lead Acquisition", "Prospective students review dynamic curriculum breakdowns and submit direct WhatsApp inquiries captured automatically into the Admin CRM."),
        ("2. Account Creation & Batch Selection", "Students create an operative account with encrypted credentials and select their targeted training cohort."),
        ("3. Bank Transfer & Receipt Upload", "Students submit bank payment receipts using the in-browser canvas-compressed image uploader."),
        ("4. Admin 1-Click Verification", "Staff inspect receipt details in Command HQ and execute a 1-click verification action."),
        ("5. Automated Batch Enrolment", "The database automatically deducts an available seat and unlocks the enrolled program directive on the student's profile."),
        ("6. LMS Video Vault & Zoom Launch", "Student dashboard immediately transitions to active training status, unlocking live Zoom Masterminds and gated lessons.")
    ]

    for s_title, s_desc in flow_steps:
        p = doc.add_paragraph()
        r1 = p.add_run(f"▶ {s_title}\n")
        r1.font.bold = True
        r1.font.color.rgb = C_GOLD
        r1.font.size = Pt(10.5)
        r2 = p.add_run(s_desc)
        r2.font.color.rgb = C_CHARCOAL
        r2.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(8)

    doc.add_heading("Technical Topology & Quality Assurance", level=2)

    qa_table = doc.add_table(rows=6, cols=2)
    qa_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(qa_table)

    qa_data = [
        ("Frontend Technology", "React 19, TypeScript, Vite, Vanilla CSS, Framer Motion 13, Lucide"),
        ("Backend Technology", "Node.js (ESM), Express 5, esbuild Serverless Bundle, Helmet, Zod"),
        ("Database Infrastructure", "PostgreSQL (Supabase) via Prisma ORM 6.19 with PgBouncer Pooling"),
        ("Cloud Deployment Target", "Vercel Edge & Serverless Functions (Single unified mono-repo build)"),
        ("Automated Test Suite", "Vitest — 32/32 Passed (100% Pass Rate across Security & Phase 7 APIs)"),
        ("API Documentation", "OpenAPI 3.0 / Swagger UI integrated at /api/docs")
    ]

    for idx, (label, val) in enumerate(qa_data):
        cells = qa_table.rows[idx].cells
        cells[0].text = label
        cells[1].text = val
        set_cell_background(cells[0], "F1F5F9")
        cells[0].paragraphs[0].runs[0].font.bold = True
        cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
        cells[1].paragraphs[0].runs[0].font.size = Pt(9.5)
        cells[0].width = Inches(2.2)
        cells[1].width = Inches(4.7)

    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    # Executive Signoff Block
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_cells = sign_table.rows[0].cells
    sign_cells[0].width = Inches(3.45)
    sign_cells[1].width = Inches(3.45)
    
    p_s1 = sign_cells[0].paragraphs[0]
    p_s1.add_run("PROPOSAL PREPARED BY:\n").font.bold = True
    p_s1.add_run("Unity Warriors Empire Technical Team\nSystem Architecture & Engineering Division")
    p_s1.runs[0].font.size = Pt(9.5)
    p_s1.runs[0].font.color.rgb = C_PRIMARY = C_NAVY
    
    p_s2 = sign_cells[1].paragraphs[0]
    p_s2.add_run("DEPLOYMENT STATUS:\n").font.bold = True
    p_s2.add_run("Production-Ready (Passed 32/32 Tests)\nStatus: Live on Vercel Edge Cloud")
    p_s2.runs[0].font.size = Pt(9.5)
    p_s2.runs[0].font.color.rgb = C_GREEN

    # Save Output
    output_path = "/Users/madushanminuwantha/Desktop/project/UWE/UWE_Commercial_Business_Proposal.docx"
    doc.save(output_path)
    print(f"Premium Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_premium_business_proposal()
